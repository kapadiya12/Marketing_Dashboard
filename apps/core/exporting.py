"""
Shared table export (CSV / Excel / PDF).

A *table* is ``{"title": str, "headers": [...], "rows": [[...], ...]}``. Views build
tables from permission-scoped data and call ``export_tables`` when ``?format=`` is
csv, xlsx or pdf.
"""
import csv
import datetime
import io
from decimal import Decimal

from django.http import HttpResponse
from django.utils import timezone
from django.utils.text import slugify

EXPORT_FORMATS = ("csv", "xlsx", "pdf")
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def plain(value):
    """Convert a cell value to plain exportable text/number."""
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return timezone.localtime(value).strftime("%Y-%m-%d %H:%M") if timezone.is_aware(value) else \
            value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (int, float, str)):
        return value
    if hasattr(value, "get_full_name"):
        return value.get_full_name() or value.get_username()
    return str(value)


def safe_cell(value):
    """Neutralise spreadsheet formula injection."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def requested_format(request):
    fmt = request.GET.get("format", "")
    return fmt if fmt in EXPORT_FORMATS else None


def export_filename(*parts):
    stem = "-".join(slugify(str(p)) for p in parts if p) or "export"
    return f"{stem}-{timezone.localdate():%Y%m%d}"


def csv_tables(tables, filename):
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}.csv"'
    response.write("﻿")  # BOM so Excel opens UTF-8 correctly
    writer = csv.writer(response)
    multi = len(tables) > 1
    for t in tables:
        if multi:
            writer.writerow([t["title"]])
        writer.writerow(t["headers"])
        for row in t["rows"]:
            writer.writerow([safe_cell(plain(v)) for v in row])
        if multi:
            writer.writerow([])
    return response


def xlsx_tables(tables, filename):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for t in tables:
        name = (slugify(t["title"]) or "sheet")[:28]
        base, n = name, 2
        while name in used:
            name = f"{base[:25]}-{n}"
            n += 1
        used.add(name)
        ws = wb.create_sheet(title=name)
        ws.append(list(t["headers"]))
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1E293B")
        for row in t["rows"]:
            ws.append([safe_cell(plain(v)) for v in row])
        for col in ws.columns:
            width = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
    if not tables:
        wb.create_sheet("empty").append(["No data"])
    buf = io.BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(),
                            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    return response


def _pdf_text(value):
    # The built-in PDF fonts have no rupee glyph.
    return str(plain(value)).replace("₹", "INR ")


def pdf_tables(tables, filename, title="", subtitle=""):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24,
                            bottomMargin=24, title=title or filename)
    styles = getSampleStyleSheet()
    cell_style = styles["BodyText"].clone("cell", fontSize=7.5, leading=9)
    head_style = cell_style.clone("head", textColor=colors.white, fontName="Helvetica-Bold")
    story = []
    if title:
        story.append(Paragraph(escape(_pdf_text(title)), styles["Title"]))
    story.append(Paragraph(escape(subtitle or f"Generated {timezone.localtime():%Y-%m-%d %H:%M}"),
                           styles["Normal"]))
    story.append(Spacer(1, 10))
    for t in tables:
        story.append(Paragraph(escape(_pdf_text(t["title"])), styles["Heading3"]))
        data = [[Paragraph(escape(_pdf_text(h)), head_style) for h in t["headers"]]]
        for row in t["rows"]:
            data.append([Paragraph(escape(_pdf_text(v))[:1500], cell_style) for v in row])
        if len(data) == 1:
            data.append([Paragraph("No data", cell_style)] + [""] * (len(t["headers"]) - 1))
        width = doc.width / max(len(t["headers"]), 1)
        tbl = Table(data, repeatRows=1, colWidths=[width] * len(t["headers"]))
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.extend([tbl, Spacer(1, 12)])
    doc.build(story)
    response = HttpResponse(buf.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}.pdf"'
    return response


def export_tables(tables, fmt, filename, title="", subtitle=""):
    if fmt == "csv":
        return csv_tables(tables, filename)
    if fmt == "xlsx":
        return xlsx_tables(tables, filename)
    if fmt == "pdf":
        return pdf_tables(tables, filename, title, subtitle)
    raise ValueError(f"Unsupported export format: {fmt}")
