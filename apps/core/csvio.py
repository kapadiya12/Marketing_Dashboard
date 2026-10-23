"""CSV / Excel export and validated CSV import for registry resources."""
import csv
import datetime
import io
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.db import models
from django.http import HttpResponse
from django.utils import timezone
from django.utils.text import slugify

from .forms import CF_PREFIX, custom_field_definitions
from .models import CustomFieldValue
from .statuses import get_status

User = get_user_model()

# Cells starting with these characters are prefixed to avoid spreadsheet formula injection.
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _safe_cell(value):
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def export_columns(resource):
    cols = ["id"] + [f for f in resource.form_fields] + ["created_at", "updated_at"]
    seen, out = set(), []
    for c in cols:
        if c not in seen:
            seen.add(c)
            out.append(c)
    return out


def cell_value(obj, name, resource=None):
    if name == "status" and resource is not None:
        info = get_status(resource.key, obj.status)
        return info.label if info else obj.status
    try:
        field = obj._meta.get_field(name)
    except Exception:
        field = None
    if field is not None and field.many_to_many:
        return "; ".join(str(o) for o in getattr(obj, name).all())
    if field is not None and field.choices:
        return getattr(obj, f"get_{name}_display")()
    value = getattr(obj, name, "")
    if callable(value):
        value = value()
    if isinstance(value, models.Model):
        if isinstance(value, User):
            return value.username
        if hasattr(value, "slug") and name == "product":
            return value.slug
        return str(value)
    if isinstance(value, datetime.datetime):
        return timezone.localtime(value).strftime("%Y-%m-%d %H:%M")
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, bool):
        return "yes" if value else "no"
    if value is None:
        return ""
    if hasattr(value, "url") and hasattr(value, "name"):
        return value.name or ""
    return value


def rows_for_export(resource, queryset):
    cols = export_columns(resource)
    defs = list(custom_field_definitions(resource.key))
    header = cols + [CF_PREFIX + d.key for d in defs]
    yield header
    ct_values = {}
    if defs:
        from django.contrib.contenttypes.models import ContentType

        ct = ContentType.objects.get_for_model(resource.model)
        for cv in CustomFieldValue.objects.filter(content_type=ct, definition__in=defs,
                                                  object_id__in=queryset.values("pk")):
            ct_values[(cv.object_id, cv.definition_id)] = cv.value
    for obj in queryset.iterator(chunk_size=500):
        row = [cell_value(obj, c, resource) for c in cols]
        for d in defs:
            v = ct_values.get((obj.pk, d.id), "")
            row.append("; ".join(v) if isinstance(v, list) else v)
        yield [_safe_cell(v) for v in row]


def csv_response(resource, queryset, filename=None):
    filename = filename or f"{resource.key}-{timezone.localdate():%Y%m%d}.csv"
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response.write("﻿")  # BOM so Excel opens UTF-8 correctly
    writer = csv.writer(response)
    for row in rows_for_export(resource, queryset):
        writer.writerow(row)
    return response


def xlsx_response(sheets, filename):
    """sheets: list of (title, rows_iterable)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets:
        ws = wb.create_sheet(title=slugify(title)[:31] or "Sheet")
        for i, row in enumerate(rows):
            ws.append([float(v) if isinstance(v, Decimal) else v for v in row])
            if i == 0:
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="1E293B")
        for col in ws.columns:
            width = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
    buf = io.BytesIO()
    wb.save(buf)
    response = HttpResponse(
        buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def import_template_columns(resource):
    cols = [c for c in resource.form_fields]
    defs = list(custom_field_definitions(resource.key))
    return cols + [CF_PREFIX + d.key for d in defs]
