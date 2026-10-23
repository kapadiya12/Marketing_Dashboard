"""Clickable boxes (KPI drill-down links) and exports everywhere."""
import datetime
from decimal import Decimal

import pytest
from django.urls import reverse
from django.utils import timezone

from apps.core.registry import registry

from .factories import ALL_RESOURCES, create_record

CONTENT_TYPES = {
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def _populate(product, platform, users=None):
    today = timezone.localdate()
    past, future = today - datetime.timedelta(days=3), today + datetime.timedelta(days=5)
    assignee = users["member_a"] if users else None
    create_record("listings", product, platform, status="approved", verification_status="verified")
    create_record("community", product, status="published", content_approval="pending", follow_up_date=past,
                  upvotes=5)
    rec = create_record("seo", product, target_ranking=5, measured_ranking=3, baseline_ranking=9)
    rec.ai_mentioned = True
    rec.save()
    create_record("seo", product, suffix="2", measured_ranking=20)
    create_record("advertising", product, leads=4, budget=Decimal("100"), actual_spend=Decimal("50"))
    create_record("events", product, event_date=future, registration_deadline=future, leads_collected=2)
    create_record("events", product, suffix="2", status="attended", event_date=past)
    create_record("outreach", product, delivery_status="delivered", reply_status="positive", meeting_booked=True,
                  follow_up_date=past)
    create_record("outreach", product, suffix="2", unsubscribed=True)
    create_record("email_campaigns", product, status="active")
    create_record("leads", product, status="qualified", unsubscribed=True)
    create_record("meetings", product, start_at=timezone.now() + datetime.timedelta(days=1),
                  opportunity_value=Decimal("1000"), follow_up_date=past)
    create_record("consultants", product, status="engaged", follow_up_date=past)
    create_record("investors", product, status="meeting_scheduled", follow_up_date=future)
    create_record("tasks", product, due_date=past, priority="urgent", assigned_to=assignee)
    create_record("tasks", product, suffix="2", due_date=today, assigned_to=assignee)
    create_record("tasks", product, suffix="3", status="done")


@pytest.mark.parametrize("key", ALL_RESOURCES)
def test_every_module_kpi_box_opens_matching_records(key, admin_user, products, platform, login):
    """Clicking a countable KPI box shows a list whose record count equals the box value."""
    _populate(products[0], platform)
    client = login(admin_user)
    resp = client.get(reverse("records:list", args=[key]))
    cards = resp.context["kpis"]
    assert cards, f"{key} has no KPI cards"
    for card in cards:
        assert card["url"], f"{key} KPI '{card['label']}' is not clickable"
        target = client.get(card["url"])
        assert target.status_code == 200, (key, card["label"], card["url"])
        if card["exact"]:
            assert target.context["total_count"] == card["value"], (key, card["label"], card["url"])


def test_kpi_links_keep_current_filters(admin_user, products, platform, login):
    a, b = products
    create_record("tasks", a, status="done")
    create_record("tasks", b, status="done", suffix="b")
    client = login(admin_user)
    resp = client.get(reverse("records:list", args=["tasks"]) + "?product=proceaze")
    completed = next(c for c in resp.context["kpis"] if c["label"] == "Completed")
    assert completed["value"] == 1
    assert "product=proceaze" in completed["url"] and "state=completed" in completed["url"]
    assert client.get(completed["url"]).context["total_count"] == 1


def test_preset_chip_and_unknown_preset(admin_user, products, login):
    create_record("tasks", products[0], priority="urgent")
    client = login(admin_user)
    resp = client.get(reverse("records:list", args=["tasks"]) + "?preset=urgent")
    assert resp.context["active_preset"] == "Open urgent / high-priority tasks"
    assert resp.context["total_count"] == 1
    resp = client.get(reverse("records:list", args=["tasks"]) + "?preset=bogus")
    assert resp.context["active_preset"] is None and resp.context["total_count"] == 1


@pytest.mark.parametrize("scope", ["org", "org-product", "product"])
def test_dashboard_boxes_open_matching_records(scope, admin_user, products, platform, login):
    a, _ = products
    _populate(a, platform)
    client = login(admin_user)
    if scope == "product":
        url = reverse("products:dashboard", args=[a.slug])
    else:
        url = reverse("analytics:dashboard") + ("?product=proceaze" if scope == "org-product" else "")
    resp = client.get(url)
    k = resp.context["k"]
    cards = {c["key"]: c for c in resp.context["kpi_cards"]}
    assert all(c["url"] for c in cards.values())
    for card in cards.values():
        assert client.get(card["url"]).status_code == 200, card["url"]
    for key, metric_value in [("activities", k["activities"]), ("completed", k["completed"]),
                              ("pending", k["pending"])]:
        drill = client.get(cards[key]["url"])
        assert drill.context["total"] == metric_value, (scope, key)
    assert "overdue" not in cards and "listings" not in cards and "products" not in cards  # hidden by default
    qs = "&product=proceaze" if scope != "org" else ""
    assert client.get(reverse("analytics:drilldown") + "?metric=overdue" + qs).context["total"] == k["overdue"]
    assert client.get(cards["meetings"]["url"]).context["total_count"] == k["meetings"]
    assert client.get(cards["events"]["url"]).context["total_count"] == k["events_upcoming"]
    assert client.get(cards["consultants"]["url"]).context["total_count"] == k["consultants_engaged"]
    assert client.get(cards["investors"]["url"]).context["total_count"] == k["investors_active"]


def test_module_cards_and_comparison_and_workload_links(admin_user, users, products, platform, login):
    a, b = products
    _populate(a, platform, users)
    client = login(admin_user)
    resp = client.get(reverse("products:dashboard", args=[a.slug]))
    for card in resp.context["module_cards"]:
        for stat in ("total", "completed", "pending", "overdue"):
            target = client.get(card["urls"][stat])
            assert target.context["total_count"] == card[stat], (card["module"].key, stat)
    org = client.get(reverse("analytics:dashboard"))
    for row in org.context["comparison"]:
        assert client.get(row["urls"]["activities"]).context["total"] == row["activities"]
        assert client.get(row["urls"]["overdue"]).context["total"] == row["overdue"]
        assert client.get(row["urls"]["meetings"]).context["total_count"] == row["meetings"]
    workload = org.context["workload"]
    assert workload and workload[0]["user"] == users["member_a"]
    drill = client.get(workload[0]["url"])
    assert drill.context["total"] == workload[0]["open"]


def test_dashboard_charts_are_clickable(admin_user, products, platform, login):
    _populate(products[0], platform)
    client = login(admin_user)
    charts = {c["id"]: c for c in client.get(reverse("analytics:dashboard")).context["charts"]}
    assert len(charts["activity_by_module"]["links"]) == 10
    assert charts["product_comparison"]["links"]
    for link in charts["activity_by_module"]["links"]:
        assert client.get(link).status_code == 200
    module_charts = client.get(reverse("records:list", args=["investors"])).context["charts"]
    assert all(c["links"] or c["filter"] for c in module_charts)


def test_drilldown_permissions(users, products, platform, login):
    a, b = products
    create_record("tasks", a, due_date=timezone.localdate() - datetime.timedelta(days=1))
    client = login(users["member_b"])
    resp = client.get(reverse("analytics:drilldown") + "?metric=overdue")
    assert resp.context["total"] == 0
    assert client.get(reverse("analytics:drilldown") + "?metric=overdue&product=proceaze").status_code == 404
    assert client.get(reverse("analytics:drilldown") + "?metric=nope").status_code == 404


# --------------------------------------------------------------------------- exports

def _assert_file(resp, fmt):
    assert resp.status_code == 200, resp.content[:300]
    assert resp["Content-Type"].startswith(CONTENT_TYPES[fmt])
    assert "attachment" in resp["Content-Disposition"]
    if fmt == "pdf":
        assert resp.content.startswith(b"%PDF")
    if fmt == "xlsx":
        assert resp.content.startswith(b"PK")


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "pdf"])
def test_exports_everywhere(fmt, admin_user, users, products, platform, login):
    a, _ = products
    _populate(a, platform, users)
    client = login(admin_user)
    inv = registry.get("investors").model.objects.first()
    urls = [
        reverse("analytics:dashboard"),
        reverse("analytics:dashboard") + "?product=proceaze&period=30d",
        reverse("products:dashboard", args=[a.slug]),
        reverse("analytics:drilldown") + "?metric=activities",
        reverse("analytics:drilldown") + "?metric=overdue&product=proceaze",
        reverse("analytics:drilldown") + "?metric=leads",
        reverse("records:detail_export", args=["investors", inv.pk]),
        reverse("records:export", args=["meetings"]),
        reverse("products:record_export", args=[a.slug, "tasks"]),
        reverse("tasks:overview"),
        reverse("search:global") + "?q=Example",
        reverse("admin_panel:users"),
        reverse("admin_panel:audit"),
        reverse("admin_panel:attachments"),
        reverse("admin_panel:config_list", args=["listing-platforms"]),
        reverse("admin_panel:config_list", args=["statuses"]),
        reverse("admin_panel:product_edit", args=[a.pk]),
        reverse("reports:index") + "?type=performance&period=all",
    ]
    for url in urls:
        sep = "&" if "?" in url else "?"
        _assert_file(client.get(f"{url}{sep}format={fmt}"), fmt)


def test_drilldown_excel_has_full_records(admin_user, products, platform, login):
    from io import BytesIO

    from openpyxl import load_workbook

    _populate(products[0], platform)
    client = login(admin_user)
    resp = client.get(reverse("analytics:drilldown") + "?metric=overdue&format=xlsx")
    wb = load_workbook(BytesIO(resp.content))
    assert "tasks-follow-ups" in wb.sheetnames
    assert wb["tasks-follow-ups"].max_row == 2  # header + 1 overdue task


def test_export_selected_rows_respects_permissions(users, products, login):
    a, b = products
    mine = create_record("tasks", b, title="Mine", suffix="m")
    other = create_record("tasks", a, title="Not mine", suffix="o")
    client = login(users["member_b"])
    resp = client.post(reverse("records:bulk", args=["tasks"]), {"ids": [mine.pk, other.pk], "action": "export_csv"})
    content = resp.content.decode("utf-8-sig")
    assert "Mine" in content and "Not mine" not in content


def test_viewer_can_export_selected(users, products, login):
    a, _ = products
    t = create_record("tasks", a, title="Visible")
    client = login(users["viewer_a"])
    resp = client.post(reverse("records:bulk", args=["tasks"]), {"ids": [t.pk], "action": "export_xlsx"})
    _assert_file(resp, "xlsx")
    resp = client.post(reverse("records:bulk", args=["tasks"]), {"ids": [t.pk], "action": "archive"})
    t.refresh_from_db()
    assert not t.is_archived


def test_detail_export_other_product_is_404(users, products, login):
    a, _ = products
    rec = create_record("investors", a)
    client = login(users["member_b"])
    assert client.get(reverse("records:detail_export", args=["investors", rec.pk]) + "?format=pdf").status_code == 404


def test_admin_exports_require_admin(users, login):
    client = login(users["member_a"])
    assert client.get(reverse("admin_panel:users") + "?format=csv").status_code == 403
    assert client.get(reverse("admin_panel:audit") + "?format=csv").status_code == 403


def test_investor_html_snapshot(admin_user, users, products, login):
    a, b = products
    create_record("investors", a, name="Alpha Fund", contact_email="secret@alpha.example", notes="internal only")
    create_record("investors", b, name="Beta Fund", suffix="b")
    client = login(admin_user)
    resp = client.get(reverse("investors:snapshot"))
    html = resp.content.decode()
    assert resp.status_code == 200 and "attachment" in resp["Content-Disposition"]
    assert "Alpha Fund" in html and "Beta Fund" in html
    assert "secret@alpha.example" not in html and "internal only" not in html
    private = client.get(reverse("investors:snapshot") + "?private=1").content.decode()
    assert "secret@alpha.example" in private
    filtered = client.get(reverse("investors:snapshot") + "?q=Alpha").content.decode()
    assert "Alpha Fund" in filtered and "Beta Fund" not in filtered
    # product isolation: a ShiftEaze member only gets ShiftEaze investors
    member = login(users["member_b"]).get(reverse("investors:snapshot")).content.decode()
    assert "Beta Fund" in member and "Alpha Fund" not in member
    assert login(users["member_b"]).get(reverse("investors:product_snapshot", args=["proceaze"])).status_code == 404
